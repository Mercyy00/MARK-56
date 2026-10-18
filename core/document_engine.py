"""
core/document_engine.py — Professional Document & Presentation Design Engine for JARVIS

Provides high-quality, visually polished PDF, DOCX, and PPTX generation
with curated color themes, professional typography, styled tables, callout boxes,
and bidirectional conversion (PDF <-> DOCX, PPTX -> PDF).
"""

from __future__ import annotations

import os
import sys
import json
import re
from pathlib import Path
from datetime import datetime
from typing import Optional, Dict, Any, List

# Core Gemini integration
from core import gemini

# Document libraries
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import pptx
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas


# ─────────────────────────────────────────────────────────────────────────────
#  Design Systems & Palettes
# ─────────────────────────────────────────────────────────────────────────────
PALETTES = {
    "navy": {
        "name": "Corporate Navy",
        "primary": "#1E3A8A",      # Deep Navy
        "secondary": "#2563EB",    # Royal Blue
        "accent": "#0D9488",       # Teal
        "dark": "#0F172A",         # Slate Dark
        "light": "#F8FAFC",        # Soft Off-white/slate
        "border": "#CBD5E1",       # Light Gray Border
        "highlight": "#EFF6FF",    # Tinted Blue
    },
    "editorial": {
        "name": "Editorial Crimson",
        "primary": "#881337",      # Burgundy / Crimson
        "secondary": "#BE123C",    # Rose
        "accent": "#D97706",       # Warm Amber
        "dark": "#1C1917",         # Charcoal
        "light": "#FAFAF9",        # Warm Off-white
        "border": "#E7E5E4",       # Stone Border
        "highlight": "#FFF1F2",    # Soft Rose Tint
    },
    "emerald": {
        "name": "Nature Emerald",
        "primary": "#064E3B",      # Deep Emerald
        "secondary": "#059669",    # Jade Green
        "accent": "#10B981",       # Mint Accent
        "dark": "#062E25",         # Deep Forest
        "light": "#F0FDF4",        # Soft Mint
        "border": "#A7F3D0",       # Light Green Border
        "highlight": "#ECFDF5",    # Tinted Green
    },
    "slate": {
        "name": "Modern Slate",
        "primary": "#334155",      # Slate
        "secondary": "#475569",    # Cool Gray
        "accent": "#0284C7",       # Sky Blue Accent
        "dark": "#0F172A",         # Very Dark Slate
        "light": "#F1F5F9",        # Light Slate
        "border": "#CBD5E1",       # Border
        "highlight": "#F8FAFC",    # Tint
    },
    "tech": {
        "name": "Tech Indigo",
        "primary": "#4338CA",      # Indigo
        "secondary": "#6366F1",    # Bright Violet
        "accent": "#EC4899",       # Neon Pink Accent
        "dark": "#1E1B4B",         # Deep Indigo Night
        "light": "#F5F3FF",        # Soft Violet
        "border": "#DDD6FE",       # Border
        "highlight": "#EEF2FF",    # Light Indigo
    }
}

def _get_palette(key: str) -> Dict[str, str]:
    k = (key or "navy").lower().strip()
    return PALETTES.get(k, PALETTES["navy"])

def _hex_to_rgb(hex_str: str) -> tuple[int, int, int]:
    h = hex_str.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


# ─────────────────────────────────────────────────────────────────────────────
#  PDF Canvas with Dynamic Header, Footer & Page Numbers (Page X of Y)
# ─────────────────────────────────────────────────────────────────────────────
class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas that accumulates total page count and prints running
    headers, footers, and 'Page X of Y' on every page except the cover page.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.doc_title = "Document"
        self.has_cover = True

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            # Skip header/footer on page 1 if there's a standalone cover
            if not (self.has_cover and self._pageNumber == 1):
                self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, total_pages: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Header rule & title
        self.drawString(54, 755, self.doc_title[:65])
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 747, 558, 747)

        # Footer rule & page numbering
        self.line(54, 48, 558, 48)
        self.drawString(54, 36, "Generated by JARVIS Document Studio")
        page_str = f"Page {self._pageNumber} of {total_pages}"
        self.drawRightString(558, 36, page_str)
        self.restoreState()


# ─────────────────────────────────────────────────────────────────────────────
#  PDF Generation Implementation
# ─────────────────────────────────────────────────────────────────────────────
def generate_pdf(doc_data: Dict[str, Any], output_path: Path, palette_key: str = "navy") -> Path:
    pal = _get_palette(palette_key)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    c_primary = colors.HexColor(pal["primary"])
    c_secondary = colors.HexColor(pal["secondary"])
    c_accent = colors.HexColor(pal["accent"])
    c_dark = colors.HexColor(pal["dark"])
    c_light = colors.HexColor(pal["light"])
    c_border = colors.HexColor(pal["border"])
    c_highlight = colors.HexColor(pal["highlight"])

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=28,
        leading=34,
        textColor=c_primary,
        alignment=0,
        spaceAfter=8
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=13,
        leading=18,
        textColor=colors.HexColor("#475569"),
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'Heading1Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=23,
        textColor=c_primary,
        spaceBefore=18,
        spaceAfter=8,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=c_secondary,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'BodyCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14.5,
        textColor=c_dark,
        spaceAfter=8
    )

    bullet_style = ParagraphStyle(
        'BulletCustom',
        parent=body_style,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=4
    )

    callout_style = ParagraphStyle(
        'CalloutText',
        parent=body_style,
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=14,
        textColor=c_dark
    )

    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=body_style,
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=body_style,
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=c_dark
    )

    story = []

    # Title & Metadata Section
    title = doc_data.get("title", "Untitled Document")
    subtitle = doc_data.get("subtitle", "")
    author = doc_data.get("author", "JARVIS System")
    date_str = doc_data.get("date", datetime.now().strftime("%B %d, %Y"))

    story.append(Paragraph(title, title_style))
    if subtitle:
        story.append(Paragraph(subtitle, subtitle_style))

    # Metadata badge / rule
    meta_text = f"<b>Author:</b> {author} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Date:</b> {date_str} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Theme:</b> {pal['name']}"
    meta_p = Paragraph(meta_text, ParagraphStyle('Meta', parent=body_style, fontSize=8.5, textColor=colors.HexColor("#64748B")))
    story.append(meta_p)
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=2, color=c_primary, spaceBefore=4, spaceAfter=14))

    # Executive Summary / Callout if provided
    exec_summary = doc_data.get("executive_summary") or doc_data.get("summary")
    if exec_summary:
        summary_p = Paragraph(f"<b>Key Takeaway:</b> {exec_summary}", callout_style)
        callout_tbl = Table([[summary_p]], colWidths=[504])
        callout_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), c_highlight),
            ('BOX', (0,0), (-1,-1), 0.5, c_border),
            ('LINELEFT', (0,0), (0,0), 3.5, c_primary),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LEFTPADDING', (0,0), (-1,-1), 12),
            ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ]))
        story.append(callout_tbl)
        story.append(Spacer(1, 14))

    # Process Sections
    for section in doc_data.get("sections", []):
        sec_title = section.get("title", "")
        sec_level = section.get("level", 1)
        if sec_title:
            if sec_level == 1:
                story.append(Paragraph(sec_title, h1_style))
                story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceBefore=1, spaceAfter=8))
            else:
                story.append(Paragraph(sec_title, h2_style))

        # Paragraphs
        for p_text in section.get("paragraphs", []):
            story.append(Paragraph(p_text, body_style))

        # Callouts inside section
        for callout in section.get("callouts", []):
            c_p = Paragraph(f"<b>Note:</b> {callout}", callout_style)
            c_table = Table([[c_p]], colWidths=[504])
            c_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), c_light),
                ('LINELEFT', (0,0), (0,0), 3, c_accent),
                ('BOX', (0,0), (-1,-1), 0.5, c_border),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('LEFTPADDING', (0,0), (-1,-1), 10),
                ('RIGHTPADDING', (0,0), (-1,-1), 10),
            ]))
            story.append(Spacer(1, 4))
            story.append(c_table)
            story.append(Spacer(1, 8))

        # Bullets
        for bullet in section.get("bullets", []):
            story.append(Paragraph(f"• &nbsp; {bullet}", bullet_style))
        if section.get("bullets"):
            story.append(Spacer(1, 6))

        # Tables
        sec_table = section.get("table")
        if sec_table and isinstance(sec_table, dict):
            headers = sec_table.get("headers", [])
            rows = sec_table.get("rows", [])
            if headers:
                num_cols = len(headers)
                col_w = 504 / num_cols
                tbl_data = [[Paragraph(h, table_header_style) for h in headers]]
                for r_idx, row in enumerate(rows):
                    row_cells = []
                    for c_idx in range(num_cols):
                        val = str(row[c_idx]) if c_idx < len(row) else ""
                        row_cells.append(Paragraph(val, table_cell_style))
                    tbl_data.append(row_cells)

                t = Table(tbl_data, colWidths=[col_w] * num_cols)
                t_styles = [
                    ('BACKGROUND', (0, 0), (-1, 0), c_primary),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('GRID', (0, 0), (-1, -1), 0.5, c_border),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ]
                # Alternating row colors
                for i in range(1, len(tbl_data)):
                    bg = c_light if (i % 2 == 1) else colors.white
                    t_styles.append(('BACKGROUND', (0, i), (-1, i), bg))

                t.setStyle(TableStyle(t_styles))
                story.append(Spacer(1, 6))
                story.append(t)
                story.append(Spacer(1, 10))

    # Canvas builder
    def canvas_maker(*args, **kwargs):
        nc = NumberedCanvas(*args, **kwargs)
        nc.doc_title = title
        nc.has_cover = False
        return nc

    doc.build(story, canvasmaker=canvas_maker)
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
#  DOCX Generation Implementation
# ─────────────────────────────────────────────────────────────────────────────
def generate_docx(doc_data: Dict[str, Any], output_path: Path, palette_key: str = "navy") -> Path:
    pal = _get_palette(palette_key)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    r_pri, g_pri, b_pri = _hex_to_rgb(pal["primary"])
    r_sec, g_sec, b_sec = _hex_to_rgb(pal["secondary"])
    r_dark, g_dark, b_dark = _hex_to_rgb(pal["dark"])
    hex_pri_clean = pal["primary"].lstrip("#")
    hex_highlight_clean = pal["highlight"].lstrip("#")
    hex_border_clean = pal["border"].lstrip("#")

    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = doc_data.get("title", "Document")[:60]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "JARVIS Document Studio  |  Confidential"
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    title = doc_data.get("title", "Untitled Document")
    subtitle = doc_data.get("subtitle", "")
    author = doc_data.get("author", "JARVIS System")
    date_str = doc_data.get("date", datetime.now().strftime("%B %d, %Y"))

    # Title Paragraph
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(title)
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(r_pri, g_pri, b_pri)

    # Subtitle Paragraph
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(12)
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = "Segoe UI"
        r_sub.font.size = Pt(13)
        r_sub.font.color.rgb = RGBColor(r_sec, g_sec, b_sec)

    # Metadata line
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta = p_meta.add_run(f"Author: {author}   |   Date: {date_str}   |   Theme: {pal['name']}")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Executive Summary / Callout Box
    exec_summary = doc_data.get("executive_summary") or doc_data.get("summary")
    if exec_summary:
        c_table = doc.add_table(rows=1, cols=1)
        c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = c_table.cell(0, 0)
        cell.width = Inches(6.5)

        # XML styling for shading and left accent border
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_highlight_clean}"/>')
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{hex_pri_clean}"/>'
            f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(shd)
        cell._tc.get_or_add_tcPr().append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(6)
        r_lead = cp.add_run("Key Summary: ")
        r_lead.font.bold = True
        r_lead.font.color.rgb = RGBColor(r_pri, g_pri, b_pri)
        r_text = cp.add_run(exec_summary)
        r_text.font.italic = True
        r_text.font.color.rgb = RGBColor(r_dark, g_dark, b_dark)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Sections
    for section in doc_data.get("sections", []):
        sec_title = section.get("title", "")
        sec_level = section.get("level", 1)

        if sec_title:
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(16 if sec_level == 1 else 10)
            p_h.paragraph_format.space_after = Pt(4)
            p_h.paragraph_format.keep_with_next = True
            r_h = p_h.add_run(sec_title)
            r_h.font.name = "Segoe UI"
            r_h.font.bold = True
            if sec_level == 1:
                r_h.font.size = Pt(16)
                r_h.font.color.rgb = RGBColor(r_pri, g_pri, b_pri)
            else:
                r_h.font.size = Pt(12.5)
                r_h.font.color.rgb = RGBColor(r_sec, g_sec, b_sec)

        # Paragraphs
        for p_text in section.get("paragraphs", []):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(p_text)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(r_dark, g_dark, b_dark)

        # Callouts
        for callout in section.get("callouts", []):
            c_table = doc.add_table(rows=1, cols=1)
            c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = c_table.cell(0, 0)
            cell.width = Inches(6.5)

            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_highlight_clean}"/>')
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{hex_pri_clean}"/>'
                f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
                f'</w:tcBorders>'
            )
            cell._tc.get_or_add_tcPr().append(shd)
            cell._tc.get_or_add_tcPr().append(borders)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(4)
            cp.paragraph_format.space_after = Pt(4)
            r = cp.add_run(f"Note: {callout}")
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.color.rgb = RGBColor(r_dark, g_dark, b_dark)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Bullets
        for bullet in section.get("bullets", []):
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(3)
            r = bp.add_run(bullet)
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(r_dark, g_dark, b_dark)

        # Tables
        sec_table = section.get("table")
        if sec_table and isinstance(sec_table, dict):
            headers = sec_table.get("headers", [])
            rows = sec_table.get("rows", [])
            if headers:
                num_cols = len(headers)
                t = doc.add_table(rows=len(rows) + 1, cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                col_width = Inches(6.5 / num_cols)

                # Header Row
                hdr_row = t.rows[0]
                # Repeat header row across pages
                trPr = hdr_row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

                for idx, text in enumerate(headers):
                    cell = hdr_row.cells[idx]
                    cell.width = col_width
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_pri_clean}"/>')
                    cell._tc.get_or_add_tcPr().append(shd)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(text)
                    r.font.name = "Segoe UI"
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Body Rows
                for r_idx, row_data in enumerate(rows):
                    row = t.rows[r_idx + 1]
                    bg_fill = hex_highlight_clean if (r_idx % 2 == 1) else "FFFFFF"
                    for c_idx in range(num_cols):
                        cell = row.cells[c_idx]
                        cell.width = col_width
                        if bg_fill != "FFFFFF":
                            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_fill}"/>')
                            cell._tc.get_or_add_tcPr().append(shd)
                        val = str(row_data[c_idx]) if c_idx < len(row_data) else ""
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        r = p.add_run(val)
                        r.font.name = "Calibri"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(r_dark, g_dark, b_dark)

                doc.add_paragraph().paragraph_format.space_after = Pt(8)

    doc.save(str(output_path))
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
#  PPTX Presentation Generation Implementation (16:9 Widescreen)
# ─────────────────────────────────────────────────────────────────────────────
def generate_pptx(deck_data: Dict[str, Any], output_path: Path, palette_key: str = "navy") -> Path:
    pal = _get_palette(palette_key)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    r_pri, g_pri, b_pri = _hex_to_rgb(pal["primary"])
    r_sec, g_sec, b_sec = _hex_to_rgb(pal["secondary"])
    r_dark, g_dark, b_dark = _hex_to_rgb(pal["dark"])
    r_light, g_light, b_light = _hex_to_rgb(pal["light"])
    r_border, g_border, b_border = _hex_to_rgb(pal["border"])

    prs = pptx.Presentation()
    # 16:9 widescreen format
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank_layout = prs.slide_layouts[6]

    title = deck_data.get("title", "Executive Presentation")
    subtitle = deck_data.get("subtitle", "")
    author = deck_data.get("author", "JARVIS System")
    date_str = deck_data.get("date", datetime.now().strftime("%B %d, %Y"))

    # Slide 1: Cover Slide
    s1 = prs.slides.add_slide(blank_layout)

    # Accent decorative background band
    band = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(0.4), PptInches(7.5))
    band.fill.solid()
    band.fill.fore_color.rgb = PptRGBColor(r_pri, g_pri, b_pri)
    band.line.fill.background()

    # Cover card container
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(1.2), PptInches(1.5), PptInches(10.9), PptInches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = PptRGBColor(r_light, g_light, b_light)
    card.line.color.rgb = PptRGBColor(r_border, g_border, b_border)
    card.line.width = PptPt(1)

    # Title inside card
    tx_box = s1.shapes.add_textbox(PptInches(1.8), PptInches(2.2), PptInches(9.7), PptInches(2.2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.name = "Segoe UI"
    p_t.font.size = PptPt(40)
    p_t.font.bold = True
    p_t.font.color.rgb = PptRGBColor(r_pri, g_pri, b_pri)

    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = PptPt(20)
        p_sub.font.color.rgb = PptRGBColor(r_sec, g_sec, b_sec)

    p_meta = tf.add_paragraph()
    p_meta.text = f"\nPresented by: {author}   |   {date_str}"
    p_meta.font.name = "Calibri"
    p_meta.font.size = PptPt(13)
    p_meta.font.color.rgb = PptRGBColor(0x64, 0x74, 0x8B)

    # Slides loop
    slides_data = deck_data.get("slides", [])
    for slide_info in slides_data:
        s = prs.slides.add_slide(blank_layout)
        s_title = slide_info.get("title", "")
        s_layout = slide_info.get("layout", "card_grid")  # card_grid, kpi_metrics, split, bullets

        # Slide Header
        h_box = s.shapes.add_textbox(PptInches(0.8), PptInches(0.6), PptInches(11.7), PptInches(1.0))
        htf = h_box.text_frame
        hp = htf.paragraphs[0]
        hp.text = s_title
        hp.font.name = "Segoe UI"
        hp.font.size = PptPt(28)
        hp.font.bold = True
        hp.font.color.rgb = PptRGBColor(r_pri, g_pri, b_pri)

        # Accent line under title
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.8), PptInches(1.5), PptInches(11.7), PptInches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = PptRGBColor(r_sec, g_sec, b_sec)
        line.line.fill.background()

        if s_layout == "kpi_metrics":
            # 3 or 4 big stat boxes
            stats = slide_info.get("stats", [])
            count = max(1, min(len(stats), 4))
            gap = 0.4
            total_w = 11.7
            box_w = (total_w - (count - 1) * gap) / count

            for i, stat in enumerate(stats[:4]):
                left = 0.8 + i * (box_w + gap)
                box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(left), PptInches(2.3), PptInches(box_w), PptInches(3.8))
                box.fill.solid()
                box.fill.fore_color.rgb = PptRGBColor(r_light, g_light, b_light)
                box.line.color.rgb = PptRGBColor(r_border, g_border, b_border)

                tx = s.shapes.add_textbox(PptInches(left + 0.2), PptInches(2.6), PptInches(box_w - 0.4), PptInches(3.2))
                stf = tx.text_frame
                stf.word_wrap = True

                p_val = stf.paragraphs[0]
                p_val.text = stat.get("value", "")
                p_val.font.name = "Segoe UI"
                p_val.font.size = PptPt(44)
                p_val.font.bold = True
                p_val.font.color.rgb = PptRGBColor(r_pri, g_pri, b_pri)
                p_val.alignment = PP_ALIGN.CENTER

                p_lbl = stf.add_paragraph()
                p_lbl.text = stat.get("label", "")
                p_lbl.font.name = "Segoe UI"
                p_lbl.font.size = PptPt(16)
                p_lbl.font.bold = True
                p_lbl.font.color.rgb = PptRGBColor(r_sec, g_sec, b_sec)
                p_lbl.alignment = PP_ALIGN.CENTER

                p_desc = stf.add_paragraph()
                p_desc.text = f"\n{stat.get('description', '')}"
                p_desc.font.name = "Calibri"
                p_desc.font.size = PptPt(12)
                p_desc.font.color.rgb = PptRGBColor(r_dark, g_dark, b_dark)
                p_desc.alignment = PP_ALIGN.CENTER

        elif s_layout == "split":
            # 2 Columns (Comparison, Problem/Solution, or Overview)
            col1 = slide_info.get("col1", {})
            col2 = slide_info.get("col2", {})
            for idx, col_data in enumerate([col1, col2]):
                left = 0.8 + idx * 6.0
                card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(left), PptInches(2.1), PptInches(5.7), PptInches(4.4))
                card.fill.solid()
                card.fill.fore_color.rgb = PptRGBColor(r_light, g_light, b_light)
                card.line.color.rgb = PptRGBColor(r_border, g_border, b_border)

                tx = s.shapes.add_textbox(PptInches(left + 0.3), PptInches(2.4), PptInches(5.1), PptInches(3.8))
                ctf = tx.text_frame
                ctf.word_wrap = True

                p_head = ctf.paragraphs[0]
                p_head.text = col_data.get("heading", "Section")
                p_head.font.name = "Segoe UI"
                p_head.font.size = PptPt(20)
                p_head.font.bold = True
                p_head.font.color.rgb = PptRGBColor(r_pri, g_pri, b_pri)

                for item in col_data.get("items", []):
                    p_it = ctf.add_paragraph()
                    p_it.text = f"•  {item}"
                    p_it.font.name = "Calibri"
                    p_it.font.size = PptPt(14)
                    p_it.font.color.rgb = PptRGBColor(r_dark, g_dark, b_dark)

        elif s_layout == "card_grid":
            # 3 Container Cards (Standard executive layout)
            cards = slide_info.get("cards", [])
            count = max(1, min(len(cards), 3))
            gap = 0.5
            box_w = (11.7 - (count - 1) * gap) / count

            for i, card_info in enumerate(cards[:3]):
                left = 0.8 + i * (box_w + gap)
                card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(left), PptInches(2.1), PptInches(box_w), PptInches(4.4))
                card.fill.solid()
                card.fill.fore_color.rgb = PptRGBColor(r_light, g_light, b_light)
                card.line.color.rgb = PptRGBColor(r_border, g_border, b_border)

                tx = s.shapes.add_textbox(PptInches(left + 0.3), PptInches(2.4), PptInches(box_w - 0.6), PptInches(3.8))
                ctf = tx.text_frame
                ctf.word_wrap = True

                p_head = ctf.paragraphs[0]
                p_head.text = card_info.get("title", "")
                p_head.font.name = "Segoe UI"
                p_head.font.size = PptPt(18)
                p_head.font.bold = True
                p_head.font.color.rgb = PptRGBColor(r_pri, g_pri, b_pri)

                p_sub = ctf.add_paragraph()
                p_sub.text = f"{card_info.get('description', '')}\n"
                p_sub.font.name = "Calibri"
                p_sub.font.size = PptPt(13)
                p_sub.font.color.rgb = PptRGBColor(r_dark, g_dark, b_dark)

                for bullet in card_info.get("bullets", []):
                    p_b = ctf.add_paragraph()
                    p_b.text = f"•  {bullet}"
                    p_b.font.name = "Calibri"
                    p_b.font.size = PptPt(12)
                    p_b.font.color.rgb = PptRGBColor(r_sec, g_sec, b_sec)

        else: # Default bullet / narrative slide
            tx = s.shapes.add_textbox(PptInches(1.0), PptInches(2.0), PptInches(11.3), PptInches(4.6))
            btf = tx.text_frame
            btf.word_wrap = True

            body_text = slide_info.get("body", "")
            if body_text:
                p_b = btf.paragraphs[0]
                p_b.text = body_text
                p_b.font.name = "Calibri"
                p_b.font.size = PptPt(16)
                p_b.font.color.rgb = PptRGBColor(r_dark, g_dark, b_dark)

            bullets = slide_info.get("bullets", [])
            for bullet in bullets:
                p_it = btf.add_paragraph()
                p_it.text = f"•  {bullet}"
                p_it.font.name = "Calibri"
                p_it.font.size = PptPt(16)
                p_it.font.color.rgb = PptRGBColor(r_dark, g_dark, b_dark)

    prs.save(str(output_path))
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
#  File Conversion Engine (Bidirectional PDF <-> DOCX, PPTX -> PDF)
# ─────────────────────────────────────────────────────────────────────────────
def convert_pdf_to_docx(pdf_path: str | Path, docx_path: str | Path) -> str:
    """High-fidelity PDF to DOCX conversion using pdf2docx with PyMuPDF layout engine."""
    src = Path(pdf_path).resolve()
    dst = Path(docx_path).resolve()
    if not src.exists():
        return f"Error: Source PDF not found: {src}"

    dst.parent.mkdir(parents=True, exist_ok=True)

    try:
        from pdf2docx import Converter
        cv = Converter(str(src))
        cv.convert(str(dst), start=0, end=None)
        cv.close()
        return f"Successfully converted PDF to DOCX: {dst.name} ({dst.stat().st_size:,} bytes)"
    except Exception as e:
        # Fallback using pdfplumber + python-docx
        try:
            import pdfplumber
            doc = docx.Document()
            doc.add_heading(src.stem, level=0)
            with pdfplumber.open(src) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text:
                        for para in text.split("\n\n"):
                            if para.strip():
                                doc.add_paragraph(para.strip())
            doc.save(str(dst))
            return f"Converted PDF to DOCX using fallback text extractor: {dst.name}"
        except Exception as e2:
            return f"PDF to DOCX conversion failed: {e} (Fallback also failed: {e2})"


def convert_docx_to_pdf(docx_path: str | Path, pdf_path: str | Path) -> str:
    """Pixel-perfect DOCX to PDF conversion using Word COM on Windows, with headless fallback."""
    src = Path(docx_path).resolve()
    dst = Path(pdf_path).resolve()
    if not src.exists():
        return f"Error: Source DOCX not found: {src}"

    dst.parent.mkdir(parents=True, exist_ok=True)

    # Primary: Native Word COM on Windows
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(src), ReadOnly=True)
            doc.SaveAs(str(dst), FileFormat=17) # 17 = wdFormatPDF
            doc.Close(False)
            return f"Successfully converted DOCX to PDF via native Word engine: {dst.name} ({dst.stat().st_size:,} bytes)"
        finally:
            word.Quit()
            pythoncom.CoUninitialize()
    except Exception as e:
        # Headless Fallback: Read DOCX paragraphs and build with ReportLab
        try:
            doc = docx.Document(str(src))
            pdf_data = {
                "title": src.stem,
                "sections": [{"title": "", "paragraphs": [p.text for p in doc.paragraphs if p.text.strip()]}]
            }
            generate_pdf(pdf_data, dst)
            return f"Converted DOCX to PDF using ReportLab layout fallback: {dst.name}"
        except Exception as e2:
            return f"DOCX to PDF conversion failed: {e} (Fallback failed: {e2})"


def convert_pptx_to_pdf(pptx_path: str | Path, pdf_path: str | Path) -> str:
    """Presentation to PDF export using PowerPoint COM on Windows."""
    src = Path(pptx_path).resolve()
    dst = Path(pdf_path).resolve()
    if not src.exists():
        return f"Error: Source PPTX not found: {src}"

    dst.parent.mkdir(parents=True, exist_ok=True)

    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        ppt = win32com.client.DispatchEx('PowerPoint.Application')
        try:
            pres = ppt.Presentations.Open(str(src), WithWindow=False)
            pres.SaveAs(str(dst), 32) # 32 = ppSaveAsPDF
            pres.Close()
            return f"Successfully exported Presentation to PDF: {dst.name} ({dst.stat().st_size:,} bytes)"
        finally:
            ppt.Quit()
            pythoncom.CoUninitialize()
    except Exception as e:
        return f"PPTX to PDF conversion failed: {e}"


# ─────────────────────────────────────────────────────────────────────────────
#  AI Content Creation & Structuring via core.gemini
# ─────────────────────────────────────────────────────────────────────────────
DOC_SCHEMA = """
Return ONLY a valid JSON object matching this schema:
{
  "title": "Document Title",
  "subtitle": "Clear executive subtitle",
  "author": "JARVIS AI Assistant",
  "executive_summary": "1-3 sentences summarizing key findings or core proposal.",
  "sections": [
    {
      "title": "Section Title",
      "level": 1,
      "paragraphs": [
        "Thorough, well-written paragraph explaining the topic.",
        "Second detailed paragraph."
      ],
      "callouts": ["Important takeaway or notable milestone"],
      "bullets": ["Key bullet item 1", "Key bullet item 2"],
      "table": {
        "headers": ["Metric", "Q1 Actual", "Q2 Target", "Status"],
        "rows": [
          ["Revenue Growth", "12.4%", "15.0%", "On Track"],
          ["System Uptime", "99.98%", "99.95%", "Exceeded"]
        ]
      }
    }
  ]
}
"""

PRESENTATION_SCHEMA = """
Return ONLY a valid JSON object matching this schema:
{
  "title": "Presentation Title",
  "subtitle": "Subtitle or theme",
  "author": "JARVIS Executive Presenter",
  "slides": [
    {
      "title": "Strategic Highlights",
      "layout": "card_grid",
      "cards": [
        {
          "title": "Architecture Overhaul",
          "description": "Modernizing legacy pipelines into microservices.",
          "bullets": ["Zero downtime migration", "10x throughput"]
        },
        {
          "title": "AI Integration",
          "description": "Autonomous agents accelerating turnaround times.",
          "bullets": ["Context-aware reasoning", "Proactive assistance"]
        },
        {
          "title": "Global Deployment",
          "description": "Multi-region redundancy across tier-1 infrastructure.",
          "bullets": ["Sub-50ms latency", "Automated failover"]
        }
      ]
    },
    {
      "title": "Key Performance Metrics",
      "layout": "kpi_metrics",
      "stats": [
        {"value": "99.98%", "label": "Reliability SLA", "description": "Continuous uninterrupted service."},
        {"value": "4.2x", "label": "Speedup Factor", "description": "Accelerated turnaround for batch analytics."},
        {"value": "-35%", "label": "Cost Optimization", "description": "Reduced cloud compute expenditure."}
      ]
    },
    {
      "title": "Current State vs Target Future",
      "layout": "split",
      "col1": {
        "heading": "Current Bottlenecks",
        "items": ["Manual document drafting", "Inconsistent formatting", "Slow turnarounds"]
      },
      "col2": {
        "heading": "Automated Vision",
        "items": ["Instant high-fidelity styling", "Unified corporate branding", "One-click conversions"]
      }
    },
    {
      "title": "Next Steps & Roadmap",
      "layout": "bullets",
      "body": "Phased implementation to guarantee smooth operational transition.",
      "bullets": [
        "Deploy v1.0 across internal developer workstations",
        "Conduct user feedback survey in week 2",
        "Roll out automated compliance checks for enterprise reporting"
      ]
    }
  ]
}
"""

def generate_document_plan(prompt: str, doc_type: str = "pdf") -> Dict[str, Any]:
    """Uses core.gemini to plan a comprehensive document structure."""
    is_presentation = doc_type in ("presentation", "pptx", "slides")
    schema = PRESENTATION_SCHEMA if is_presentation else DOC_SCHEMA

    user_instruction = (
        f"You are the world's elite corporate document and presentation designer.\n"
        f"Generate comprehensive, realistic, highly professional content based on this request:\n"
        f"REQUEST: {prompt}\n\n"
        f"{schema}\n"
        f"Important: Ensure all content is substantial, informative, professional, and directly matches the JSON structure."
    )

    try:
        data = gemini.as_json(user_instruction, tier=gemini.SMART, timeout_ms=60000)
        if isinstance(data, dict):
            return data
    except Exception as e:
        print(f"[document_engine] Gemini structuring failed, falling back to minimal plan: {e}")

    # Fallback structure if Gemini is offline
    title = prompt[:50].title() if prompt else "Executive Brief"
    if is_presentation:
        return {
            "title": title,
            "subtitle": "Prepared by JARVIS",
            "slides": [
                {
                    "title": "Overview",
                    "layout": "bullets",
                    "body": prompt,
                    "bullets": ["Action item 1", "Action item 2", "Summary conclusion"]
                }
            ]
        }
    return {
        "title": title,
        "subtitle": "Generated Document",
        "executive_summary": prompt,
        "sections": [
            {
                "title": "Detailed Analysis",
                "level": 1,
                "paragraphs": [prompt],
                "bullets": ["Primary finding", "Recommended next step"]
            }
        ]
    }
